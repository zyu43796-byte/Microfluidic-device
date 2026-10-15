"""DMF_HOST UI adapted to the current ASCII firmware. Never imports legacy SerialDriver."""
import ast
import html
import math
import sys
import time
from pathlib import Path

import serial
from serial.tools import list_ports
from PyQt5.QtCore import Qt, QEvent, QTimer
from PyQt5.QtGui import QIcon, QIntValidator, QKeySequence
from PyQt5.QtWidgets import (QApplication, QMainWindow, QPushButton, QLabel, QCheckBox,
                             QSpinBox, QComboBox, QTableWidgetItem, QFileDialog,
                             QLineEdit, QAbstractItemView, QMessageBox, QWidget,
                             QVBoxLayout, QHBoxLayout, QGridLayout, QSlider, QTabWidget)
from main_window.ui2 import UI
from chip.DMF_Chip import DMF_Chip
from control import Controller, bitmap, move_ids
from lighting import OPTIONS, COLORS
from magnet_panel import MagnetPanel
from utils.path_helper import get_resource_path
from modern_workspace import install_workspace, update_summary, notify


class MainWindow(QMainWindow, UI):
    def __init__(self):
        super().__init__()
        self.SetUpUi(self)
        self.setWindowTitle('DMF 上位机 · AC 开关 / 持续驱动 / 灯光 / 磁杆')
        self.setWindowIcon(QIcon(get_resource_path('icon/dmf.png')))
        self.camera_logic = None
        self._closing = False
        self._auto_rows = []
        self._last_pinned = []
        self.scene = DMF_Chip(get_resource_path('chip/chip_point.csv'))
        self.dmf_chip_scene = self.scene
        self.graphicsView1.setScene(self.scene)
        self.control = Controller(emit=self.on_control_event)
        self._build_controls()
        install_workspace(self)
        self._bind_controls()
        self.flush_ports()
        self.log_message('先核对电极编号与实物方向，连接后点击开启 AC 模式。颜色表示选择，不是液滴位置或实测电压。')
        self.poll_timer = QTimer(self)
        self.poll_timer.setInterval(10)
        self.poll_timer.timeout.connect(self.poll)
        self.poll_timer.start()
        QApplication.instance().installEventFilter(self)
        self.refresh_enabled()
        QTimer.singleShot(0, self.fit_chip)

    def _build_controls(self):
        self.button1.setText('刷新压实联锁')
        self.button2.setText('位置开关')
        self.button3.setText('综合联锁 / 旁路')
        for label in (self.label1_button, self.label2_button, self.label3_button):
            label.setText('未知')
        self.button4.setText('设定电压 / V（40–150）')
        self.lineEdit1.setText('80')
        self.lineEdit1.setValidator(QIntValidator(40, 150, self))
        self.button5.setText('AC 频率 / Hz')
        self.lineEdit2.setText('1000')
        self.lineEdit2.setReadOnly(True)
        self.button6.setText('AC / POL 延时 / μs')
        self.lineEdit3_1.setText('40')
        self.lineEdit3_1.setReadOnly(True)
        self.button7.hide()
        self.lineEdit3_2.hide()
        self.mapping_check = QCheckBox('已核对界面电极编号与实物对应')
        self.mapping_check.setToolTip('沿用学长的电极布局；例如 1 的左邻是 2，右邻是 8。')
        self.grid_layout2.addWidget(self.mapping_check, 4, 0, 1, 7)
        self.ac_on_button = QPushButton('开启 AC 模式')
        self.ac_on_button.setToolTip('检查联锁并应用参数；选择电极后持续驱动，无60秒总运行限制。')
        self.ac_off_button = QPushButton('关闭 AC')
        self.ac_off_button.setToolTip('停止高压驱动并退出AC模式；当前固件会同时关闭灯带。升压供电需手动关闭。')
        self.grid_layout2.addWidget(self.ac_on_button, 5, 0, 1, 4)
        self.grid_layout2.addWidget(self.ac_off_button, 5, 4, 1, 3)
        self.stop_button = QPushButton('停止全部输出（Esc）')
        self.stop_button.setStyleSheet('background:#b91c1c;color:white;font-weight:bold;padding:7px;')
        self.grid_layout2.addWidget(self.stop_button, 6, 0, 1, 7)
        self.state_label = QLabel('未连接 / AC 模式已关闭')
        self.state_label.setWordWrap(True)
        self.grid_layout2.addWidget(self.state_label, 7, 0, 1, 7)
        self.output_label = QLabel('AC 输出未启动')
        self.output_label.setWordWrap(True)
        self.layout_V_Middle_V1.addWidget(self.output_label)
        self.repeat_button = QPushButton('重新下发当前连续驱动')
        self.repeat_button.hide()
        self.clear_button = QPushButton('取消全部电极（含固定）')
        self.grid_layout3.addWidget(self.repeat_button, 4, 0, 1, 5)
        self.grid_layout3.addWidget(self.clear_button, 5, 0, 1, 5)
        self.tabWidget.setTabEnabled(2, False)
        self.tabWidget.setTabText(2, '液滴检测（未适配）')
        self.tabWidget2.setEnabled(False)
        self.tabWidget2.hide()
        self.groupBox5.setTitle('外设控制 · 灯光 / 磁杆')
        self._build_light_panel()
        self.layout_V_Right_V2.removeWidget(self.light_panel)
        self.tool_tabs = QTabWidget(self.groupBox5)
        self.tool_tabs.addTab(self.light_panel, '灯光')
        self.magnet_panel = MagnetPanel(self.apply_magnet, self.stop_output)
        self.tool_tabs.addTab(self.magnet_panel, '磁杆')
        self.layout_V_Right_V2.addWidget(self.tool_tabs)
        self.tab2_tab.setHorizontalHeaderLabels(['电极编号列表', '连续保持时间(s)', '相机预览'])
        self.tab2_tab.setToolTip('每帧连续 AC 保持指定时间；切帧先 OFF。[] 为全关等待。')
        self.tab2_tab.setDragDropMode(QAbstractItemView.NoDragDrop)
        self.camera_index = QSpinBox()
        self.camera_index.setRange(0, 9)
        self.camera_index.setPrefix('相机 ')
        self.layout_V_Right_V1_H1.insertWidget(0, self.camera_index)
        self.button_close_camera.setEnabled(False)
        self.button_take_picture.setEnabled(False)
        self.button_record_picture.setEnabled(False)
        self.button_record_picture.setText('录像（原程序未实现）')
        self.groupBox4.setTitle('摄像头预览（请在开启 AC 前打开）')
        self.log_information.document().setMaximumBlockCount(1000)
        self.serial_information.document().setMaximumBlockCount(1500)

    def _build_light_panel(self):
        self.light_panel = QWidget(self.groupBox5)
        layout = QVBoxLayout(self.light_panel)
        layout.setSpacing(12)
        tip = QLabel('选择颜色或灯效，调整亮度后点击“应用灯光”。')
        tip.setWordWrap(True)
        layout.addWidget(tip)
        colors_row = QHBoxLayout()
        self.light_color_buttons = []
        for name, key, color, ink in [('白', 'white', '#f1f5f9', '#1e293b'),
                                      ('红', 'red', '#dc2626', 'white'),
                                      ('绿', 'green', '#16a34a', 'white'),
                                      ('蓝', 'blue', '#2563eb', 'white'),
                                      ('黄', 'yellow', '#facc15', '#1e293b'),
                                      ('青', 'cyan', '#06b6d4', '#1e293b'),
                                      ('紫', 'magenta', '#c026d3', 'white')]:
            button = QPushButton(name)
            button.setMinimumHeight(32)
            button.setStyleSheet(f'QPushButton {{background:{color};color:{ink};border:1px solid #94a3b8;border-radius:5px;}}'
                                'QPushButton:disabled {background:#e2e8f0;color:#94a3b8;}')
            button.setToolTip('选择静态颜色；点击应用灯光后发送')
            button.clicked.connect(lambda checked=False, value=key: self.light_choice.setCurrentIndex(self.light_choice.findData(value)))
            colors_row.addWidget(button)
            self.light_color_buttons.append(button)
        layout.addLayout(colors_row)
        grid = QGridLayout()
        grid.setVerticalSpacing(12)
        self.light_choice = QComboBox()
        for label, mode in OPTIONS:
            self.light_choice.addItem(label, mode)
        grid.addWidget(QLabel('灯光模式'), 0, 0)
        grid.addWidget(self.light_choice, 0, 1, 1, 2)
        self.light_brightness = QSpinBox()
        self.light_brightness.setRange(0, 255)
        self.light_brightness.setValue(80)
        self.light_slider = QSlider(Qt.Horizontal)
        self.light_slider.setRange(0, 255)
        self.light_slider.setValue(80)
        self.light_slider.valueChanged.connect(self.light_brightness.setValue)
        self.light_brightness.valueChanged.connect(self.light_slider.setValue)
        grid.addWidget(QLabel('亮度（0–255）'), 1, 0)
        grid.addWidget(self.light_slider, 1, 1)
        grid.addWidget(self.light_brightness, 1, 2)
        self.light_interval = QSpinBox()
        self.light_interval.setRange(10, 1000)
        self.light_interval.setValue(30)
        self.light_interval.setSuffix(' ms')
        grid.addWidget(QLabel('灯效帧间隔'), 2, 0)
        grid.addWidget(self.light_interval, 2, 1)
        grid.addWidget(QLabel('数值越小，动画越快'), 2, 2)
        grid.setColumnStretch(1, 1)
        layout.addLayout(grid)
        self.light_hint = QLabel()
        self.light_hint.setWordWrap(True)
        self.light_hint.setStyleSheet('color:#475569;')
        layout.addWidget(self.light_hint)
        actions = QHBoxLayout()
        self.light_apply_button = QPushButton('应用灯光')
        self.light_off_button = QPushButton('关闭灯带')
        self.light_apply_button.setStyleSheet('background:#0f766e;color:white;padding:8px;')
        self.light_off_button.setStyleSheet('padding:8px;')
        actions.addWidget(self.light_apply_button, 2)
        actions.addWidget(self.light_off_button, 1)
        layout.addLayout(actions)
        self.light_status = QLabel('串口未连接；请连接后控制灯光。')
        self.light_status.setWordWrap(True)
        self.light_status.setMinimumHeight(48)
        layout.addWidget(self.light_status)
        note = QLabel('灯光可在 AC 关闭时使用。关闭 AC / 停止全部输出也会关闭灯带。')
        note.setWordWrap(True)
        note.setStyleSheet('color:#64748b;')
        layout.addWidget(note)
        layout.addStretch()
        self.layout_V_Right_V2.addWidget(self.light_panel)
        self.light_choice.currentIndexChanged.connect(self.light_mode_changed)
        self.light_apply_button.clicked.connect(self.apply_light)
        self.light_off_button.clicked.connect(lambda: self.apply_light(off=True))
        self.light_mode_changed()

    def light_mode_changed(self):
        mode = self.light_choice.currentData()
        hints = {'rainbow': '彩虹颜色沿灯带连续流动。', 'meteor': '流星光头带变色拖尾；可将亮度调到112。',
                 'chase': '三个不同颜色的光头沿灯带追逐。', 'stack': '光点逐颗累加，填满后渐灭重来。',
                 'sparkle': '彩色星点分散闪烁。', 'collide': '两端彗星相向对撞，释放彩虹光波。',
                 'demo': '六种灯效轮流播放，每种10秒；统一使用所选亮度。'}
        self.light_hint.setText(hints.get(mode, '静态常亮；亮度0为熄灭，255为最高值。'))
        self.light_interval.setEnabled(self.control.ready and not self.control.busy and mode not in COLORS)

    def apply_light(self, checked=False, off=False):
        try:
            mode = 'off' if off else self.light_choice.currentData()
            self.control.set_light(mode, self.light_brightness.value(), self.light_interval.value())
            self.light_status.setText('正在发送灯光指令，等待确认…')
        except (ValueError, RuntimeError) as exc:
            self.light_status.setText(str(exc))
            self.log_message(exc)
        self.refresh_enabled()

    def _bind_controls(self):
        self.button_open_serial_port.clicked.connect(self.open_port)
        self.button_close_serial_port.clicked.connect(lambda: self.control.stop(disconnect=True))
        self.button_flush_serial_port.clicked.connect(self.flush_ports)
        for button in (self.button1, self.button2, self.button3):
            button.clicked.connect(self.request_status)
        self.ac_on_button.clicked.connect(self.arm)
        self.ac_off_button.clicked.connect(self.stop_output)
        self.stop_button.clicked.connect(self.stop_output)
        self.mapping_check.toggled.connect(self.mapping_changed)
        self.lineEdit1.textEdited.connect(self.parameters_changed)
        self.scene.serialDataSignal.connect(self.send_selection)
        for button, dr, dc in [(self.button_up, -1, 0), (self.button_down, 1, 0),
                               (self.button_left, 0, -1), (self.button_right, 0, 1)]:
            button.setShortcut(QKeySequence())
            button.clicked.connect(lambda checked=False, r=dr, c=dc: self.move_electrodes(r, c))
        self.button_zoomout.clicked.connect(lambda: self.graphicsView1.scale(.8, .8))
        self.button_zoomin.clicked.connect(lambda: self.graphicsView1.scale(1.25, 1.25))
        self.repeat_button.clicked.connect(self.send_selection)
        self.clear_button.clicked.connect(self.clear_selection)
        self.button_add.clicked.connect(self.add_current_row)
        self.button_del.clicked.connect(self.delete_row)
        self.button_sta.clicked.connect(self.start_sequence)
        self.button_end.clicked.connect(self.stop_output)
        self.button_save_data.clicked.connect(self.save_rows)
        self.button_load_data.clicked.connect(self.load_rows)
        self.button_open_camera.clicked.connect(self.open_camera)

    def apply_magnet(self, mode):
        try:
            self.control.set_magnet(mode, self.magnet_panel.duration.value())
            self.magnet_panel.status.setText('正在检查 / 下发磁杆指令，等待确认…')
        except (ValueError, RuntimeError) as exc:
            self.magnet_panel.status.setText(str(exc))
            self.log_message(exc)
        self.refresh_enabled()

    def log_message(self, message, error=False):
        if error or isinstance(message, Exception):
            notify(self, '操作失败：' + str(message), error=True)
        self.log_information.append(f'[{time.strftime("%H:%M:%S")}] {html.escape(str(message))}')

    def debug(self, record, message):
        # CameraLogic's logger interface; all messages stay on the Qt main thread.
        self.log_message(message)

    def flush_ports(self):
        current = self.comboBox.currentData()
        self.comboBox.clear()
        for port in list_ports.comports():
            self.comboBox.addItem(f'{port.device} — {port.description}', port.device)
        index = self.comboBox.findData(current)
        if index >= 0:
            self.comboBox.setCurrentIndex(index)

    def open_port(self):
        try:
            port_name = self.comboBox.currentData()
            if not port_name:
                raise ValueError('未找到串口，请连接 USB 串口并刷新')
            port = serial.Serial(port_name, 115200, timeout=0, write_timeout=.1,
                                 bytesize=8, parity='N', stopbits=1)
            try:
                self.control.connect(port)
            except Exception:
                port.close()
                raise
        except (OSError, ValueError, RuntimeError) as exc:
            self.log_message(exc)

    def request_status(self):
        try:
            self.control.request_status()
        except (ValueError, RuntimeError) as exc:
            self.log_message(exc)

    def mapping_changed(self, checked):
        if not checked and self.control.armed:
            self.stop_output()
        self.refresh_enabled()

    def parameters_changed(self):
        if self.control.armed:
            self.log_message('电压输入已修改：已请求停止，修改完成后重新点击开启 AC 模式。')
            self.stop_output()

    def arm(self):
        try:
            if not self.mapping_check.isChecked():
                raise ValueError('请先核对电极映射并勾选确认')
            self.control.arm(int(self.lineEdit1.text()))
        except (ValueError, RuntimeError) as exc:
            self.log_message(exc)

    def set_scene(self, ids, pinned=()):
        ids, pinned = set(ids), set(pinned)
        self.scene._cancel_ctrl_right_pending()
        for e in self.scene._iter_electrodes():
            e.setPinned(False)
            e.setDriveStatus(e.ID in ids)
            e.setSelected(e.ID in ids)
            if e.ID in pinned and e.ID in ids:
                e.setPinned(True)

    def send_selection(self):
        try:
            if not self.mapping_check.isChecked() or not self.control.armed:
                raise ValueError('AC 模式未开启：请先核对电极映射并点击开启 AC 模式')
            ids = self.scene.hardware_drive_ids()
            pinned = self.scene.pinned_ids()
            self.control.select_continuous(ids)
            self._last_pinned = pinned
            self.output_label.setText('正在安全切换电极；等待持续 AC 应答')
        except (ValueError, RuntimeError) as exc:
            self.log_message(exc)
            self.set_scene(self.control.ack_ids, self._last_pinned)

    def move_electrodes(self, dr, dc):
        try:
            ids = self.scene.hardware_drive_ids()
            if not ids:
                raise ValueError('请先右键选择一个起始电极')
            pinned = self.scene.pinned_ids()
            targets = move_ids(ids, pinned, dr, dc, self.scene.id_array)
            self.set_scene(targets, pinned)
            self.send_selection()
        except (ValueError, RuntimeError) as exc:
            self.log_message(exc)

    def clear_selection(self):
        self.set_scene([])
        self._last_pinned = []
        if self.control.armed:
            self.send_selection()

    def stop_output(self):
        self._auto_rows = []
        self._last_pinned = []
        self.set_scene([])
        self.control.stop()
        self.refresh_enabled()

    def add_row(self, ids=None, seconds=.1, camera='关'):
        row = self.tab2_tab.rowCount()
        self.tab2_tab.insertRow(row)
        self.tab2_tab.setItem(row, 0, QTableWidgetItem(str(ids or [])))
        self.tab2_tab.setItem(row, 1, QTableWidgetItem(str(seconds)))
        combo = QComboBox()
        combo.addItems(['关', '开'])
        combo.setCurrentText(camera)
        self.tab2_tab.setCellWidget(row, 2, combo)

    def add_current_row(self):
        ids = sorted({e.ID for e in self.scene._iter_electrodes() if e.isSelected() or e.isDrive})
        self.add_row(ids)

    def delete_row(self):
        row = self.tab2_tab.currentRow()
        self.tab2_tab.removeRow(row if row >= 0 else self.tab2_tab.rowCount() - 1)

    @staticmethod
    def parse_row(frame, seconds, camera):
        frame = str(frame).strip()
        try:
            ids = ast.literal_eval(frame) if frame.startswith('[') else [int(x.strip()) for x in frame.split(',') if x.strip()]
            if not isinstance(ids, list):
                raise ValueError('电极列表格式错误')
            bitmap(ids)
            seconds = float(seconds)
            if not math.isfinite(seconds) or not .001 <= seconds <= 10:
                raise ValueError('连续保持时间应在 0.001..10 秒')
            if camera not in ('开', '关'):
                raise ValueError('摄像头状态只能为开/关')
            return sorted(set(ids)), round(seconds * 1000), camera
        except (ValueError, SyntaxError, TypeError) as exc:
            raise ValueError(f'流程行无效：{exc}') from exc

    def read_rows(self):
        rows = []
        for row in range(self.tab2_tab.rowCount()):
            try:
                frame = self.tab2_tab.item(row, 0)
                duration = self.tab2_tab.item(row, 1)
                if frame is None or duration is None:
                    raise ValueError('电极或时间为空')
                rows.append(self.parse_row(frame.text(), duration.text(), self.tab2_tab.cellWidget(row, 2).currentText()))
            except ValueError as exc:
                raise ValueError(f'第 {row + 1} 行：{exc}') from exc
        return rows

    def start_sequence(self):
        try:
            if not self.mapping_check.isChecked():
                raise ValueError('请先确认电极映射')
            rows = self.read_rows()
            # Opening a camera during an energized step can block UI stop handling.
            if any(camera == '开' for _, _, camera in rows):
                if self.camera_logic is None or not self.camera_logic.timer_camera.isActive():
                    raise ValueError('流程包含摄像头：请先手动打开并确认图像')
            self._auto_rows = rows
            self._last_pinned = []
            self.control.sequence([(ids, duration) for ids, duration, _ in rows])
        except (ValueError, RuntimeError) as exc:
            self.log_message(exc)

    def save_rows(self):
        try:
            rows = self.read_rows()
            path, _ = QFileDialog.getSaveFileName(self, '保存流程', 'DMF流程.xlsx', 'Excel (*.xlsx)')
            if not path:
                return
            from openpyxl import Workbook
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(['帧', '驱动时间(S)', '开启摄像头'])
            for ids, ms, camera in rows:
                worksheet.append([str(ids), ms / 1000, camera])
            workbook.save(path)
            self.log_message('已保存流程：' + path)
        except (ValueError, OSError) as exc:
            self.log_message(exc)

    def load_rows(self):
        path, _ = QFileDialog.getOpenFileName(self, '加载流程', '', 'Excel (*.xlsx)')
        if not path:
            return
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = [self.parse_row(row[0], row[1], row[2]) for row in
                        workbook.active.iter_rows(min_row=2, max_col=3, values_only=True) if any(v is not None for v in row)]
            finally:
                workbook.close()
            if len(rows) > 1000:
                raise ValueError('最多允许 1000 帧')
            self.tab2_tab.setRowCount(0)
            for ids, ms, camera in rows:
                self.add_row(ids, ms / 1000, camera)
            self.log_message('已加载；请检查各帧连续保持时间。')
        except Exception as exc:
            self.log_message('加载失败，原表保留：' + str(exc), error=True)

    def open_camera(self):
        if self.control.busy or self.control.armed or self.control.magnet_active:
            self.log_message('请在开启 AC / 磁杆前打开摄像头', error=True)
            return
        if self.camera_logic is None:
            from cam.cam_module import CameraLogic
            self.button_open_camera.clicked.disconnect(self.open_camera)
            self.camera_logic = CameraLogic(self.label_cam_view, self.button_open_camera,
                                           self.button_close_camera, self.button_take_picture,
                                           None, self, self.camera_index.value())
            # Keep all subsequent opens behind the same authorization guard.
            self.button_open_camera.clicked.disconnect(self.camera_logic.open_camera)
            self.button_open_camera.clicked.connect(self.open_camera)
        self.camera_logic.open_camera()
        self.button_record_picture.setEnabled(False)

    def on_control_event(self, kind, data):
        if kind == 'error':
            notify(self, '异常：' + str(data), error=True)
        elif kind == 'connected':
            notify(self, '串口握手完成 · 请核对联锁与电极编号')
        elif kind == 'armed':
            notify(self, 'AC 模式开启已确认 · 请选择电极')
        elif kind == 'stopped':
            notify(self, str(data))
        elif kind == 'disconnected':
            notify(self, '串口已断开 · 设备输出状态无法继续确认')
        if kind in ('tx', 'rx'):
            self.serial_information.append(html.escape(('→ ' if kind == 'tx' else '← ') + data))
        elif kind == 'gpio':
            for label, key in [(self.label1_button, 'biochip_check'), (self.label2_button, 'pos_sw1')]:
                label.setText('闭合' if data[key] else '断开')
                label.setStyleSheet('color:' + ('#15803d' if data[key] else '#b91c1c'))
            valid = data['interlock_closed'] and not data['biochip_interlock_bypassed']
            self.label3_button.setText('正常' if valid else ('旁路！' if data['biochip_interlock_bypassed'] else '断开'))
            self.label3_button.setStyleSheet('color:' + ('#15803d' if valid else '#b91c1c'))
        elif kind == 'array':
            self.set_scene(data, self._last_pinned)
        elif kind == 'continuous_on':
            self.output_label.setText('连续 AC 输出中：电极 ' + ', '.join(map(str, data)) + '。注意观察 U7 温升。')
        elif kind == 'continuous_off':
            self.output_label.setText('AC 已关闭；电极全零。AC 模式仍开启，等待选择电极。')
        elif kind == 'light':
            self.light_status.setText('灯光指令已确认：' + data)
            self.log_message(data)
        elif kind == 'magnet':
            text = 'ON 指令已确认；动作保持中（无到位反馈）。' if data == 'on' else 'OFF 指令已确认；已切回默认极性（无到位反馈）。'
            self.magnet_panel.status.setText(text)
            self.log_message(text)
        elif kind == 'stopping':
            self.magnet_panel.status.setText('正在停止全部输出；磁杆状态等待确认。')
        elif kind == 'row':
            self.tab2_tab.setCurrentCell(data, 0)
            self.log_message(f'执行流程第 {data + 1} 行')
            # The camera must already be open; rows only pause/resume preview.
            if self.camera_logic and data < len(self._auto_rows):
                if self._auto_rows[data][2] == '开':
                    self.camera_logic.timer_camera.start(40)
                else:
                    self.camera_logic.timer_camera.stop()
                    self.label_cam_view.setText('相机预览暂停')
        elif kind == 'armed':
            self.set_scene([])
            self._last_pinned = []
            self.output_label.setText('AC 模式已开启；电极全零，选择电极后开始持续驱动。')
            self.log_message(f'AC 模式已开启，无60秒总运行限制：设定 {data}V / 1000Hz / 40μs。设定值不等于实测电压。')
        elif kind in ('stopped', 'error', 'disconnected', 'connected'):
            if kind == 'connected':
                self.log_message('当前 CLI 握手成功；AC 模式已关闭，灯光可单独使用。')
                self.light_status.setText('已连接；选择颜色或灯效，再点击应用。')
                self.magnet_panel.status.setText('已连接；初始化全停已确认，磁杆为默认指令状态。')
            else:
                self.log_message(data)
            if kind in ('stopped', 'error', 'disconnected'):
                self.set_scene([])
                self._last_pinned = []
                self._auto_rows = []
            if kind == 'stopped':
                self.output_label.setText(data)
                self.magnet_panel.status.setText('全停已确认；磁杆切回默认极性（不代表机械停止或到位）。')
                self.light_status.setText('已发送全部停止；六灯效版固件同时关闭灯带。')
            if kind == 'error':
                self.output_label.setText('异常：' + str(data))
                self.magnet_panel.status.setText('磁杆状态待确认：' + str(data))
                self.light_status.setText('异常，灯光/输出状态待确认：' + str(data))
            if kind == 'disconnected':
                self.magnet_panel.status.setText('串口已断开；无法继续确认磁杆状态。')
                self.light_status.setText('串口未连接；请连接后控制灯光。')
                for label in (self.label1_button, self.label2_button, self.label3_button):
                    label.setText('未知')
                    label.setStyleSheet('color:#666')

    def refresh_enabled(self):
        c = self.control
        update_summary(self)
        connected = bool(c.link and c.link.serial.is_open)
        available = c.ready and not c.busy
        self.magnet_panel.set_available(available)
        drive_ok = available and c.armed and self.mapping_check.isChecked()
        self.button_open_serial_port.setEnabled(not connected and not c.busy)
        self.button_close_serial_port.setEnabled(connected)
        self.ac_on_button.setEnabled(available and not c.armed and self.mapping_check.isChecked())
        self.ac_off_button.setEnabled(connected)
        self.stop_button.setEnabled(True)
        for widget in (self.light_choice, self.light_brightness, self.light_slider,
                       self.light_apply_button, self.light_off_button, *self.light_color_buttons):
            widget.setEnabled(available)
        self.light_interval.setEnabled(available and self.light_choice.currentData() not in COLORS)
        for b in (self.button_up, self.button_down, self.button_left, self.button_right,
                  self.repeat_button, self.button_sta):
            b.setEnabled(drive_ok)
        self.clear_button.setEnabled(not c.busy)
        self.graphicsView1.setEnabled(not c.busy)
        for widget in (self.tab2_tab, self.button_add, self.button_del, self.button_load_data,
                       self.button_save_data, self.lineEdit1, self.lineEdit3_2):
            widget.setEnabled(not c.busy)
        self.camera_index.setEnabled(self.camera_logic is None and not c.armed and not c.magnet_active and not c.busy)
        if self.camera_logic is None or not self.camera_logic.timer_camera.isActive():
            self.button_open_camera.setEnabled(not c.armed and not c.magnet_active and not c.busy)
        self.button_take_picture.setEnabled(bool(self.camera_logic and self.camera_logic.timer_camera.isActive()) and not c.armed and not c.magnet_active and not c.busy)
        status = 'AC 模式已开启（无60秒限制）' if c.armed else 'AC 模式已关闭'
        self.state_label.setText(('已连接 / ' if c.ready else '未就绪 / ') + status +
                                ('\n正在：' + c.job_name if c.busy and c.job_name not in ('AC 通信保持', '磁杆状态检查') else
                                 ('\n连续 AC 正在输出' if c.output_active else '\n当前无 AC 驱动输出')))

    def poll(self):
        self.control.tick()
        self.refresh_enabled()
        if self._closing and not self.control.busy:
            self.close()

    def fit_chip(self):
        self.graphicsView1.fitInView(self.scene.itemsBoundingRect(), Qt.KeepAspectRatio)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, 'scene'):
            QTimer.singleShot(0, self.fit_chip)

    def eventFilter(self, obj, event):
        if event.type() == QEvent.Resize and obj is self.graphicsView1:
            # Watch the view, not its viewport: zoom scrollbars resize the latter
            # and must not reset a user's zoom or cause a fit/resize feedback loop.
            QTimer.singleShot(0, self.fit_chip)
        if event.type() == QEvent.KeyPress and event.key() == Qt.Key_Escape:
            self.stop_output()
            return True
        if event.type() == QEvent.KeyRelease and event.key() == Qt.Key_Control:
            self.scene.flush_ctrl_right_pending()
        if event.type() == QEvent.KeyPress and not event.isAutoRepeat():
            focused = QApplication.focusWidget()
            if focused is self.graphicsView1 or focused is self.graphicsView1.viewport():
                directions = {Qt.Key_Left: (0, -1), Qt.Key_Right: (0, 1), Qt.Key_Up: (-1, 0), Qt.Key_Down: (1, 0)}
                if event.key() in directions:
                    self.move_electrodes(*directions[event.key()])
                    return True
        return super().eventFilter(obj, event)

    def closeEvent(self, event):
        c = self.control
        if c.link and c.link.serial.is_open:
            event.ignore()
            if not self._closing:
                self._closing = True
                c.stop(disconnect=True)
            return
        if c.uncertain and self._closing:
            QMessageBox.warning(self, '停止状态未确认', '串口停止未确认，请手动切断高压供电。')
        self.poll_timer.stop()
        QApplication.instance().removeEventFilter(self)
        if self.camera_logic:
            self.camera_logic.cleanup()
        event.accept()


def main():
    app = QApplication(sys.argv)
    from startup_animation import show_with_startup

    # Explicit QA mode follows the normal startup without opening any device.
    def on_ready(window):
        if '--self-test' not in sys.argv:
            return
        def finish():
            from cam.cam_module import CameraLogic
            from openpyxl import Workbook
            assert CameraLogic is not None and Workbook is not None
            assert window.control.link is None and window.camera_logic is None
            from control import VALID_IDS
            assert {e.ID for e in window.scene._iter_electrodes()} == VALID_IDS
            assert window.tool_tabs.count() == 2
            if '--magnet-tab' in sys.argv:
                window.tool_tabs.setCurrentIndex(1)
                app.processEvents()
            if '--screenshot' in sys.argv:
                path = sys.argv[sys.argv.index('--screenshot') + 1]
                if not window.grab().save(path):
                    app.exit(2)
                    return
            app.exit(0)
        QTimer.singleShot(500, finish)
    splash = show_with_startup(MainWindow, on_ready)
    if '--self-test' in sys.argv and '--splash-screenshot' in sys.argv:
        def capture_splash():
            path = sys.argv[sys.argv.index('--splash-screenshot') + 1]
            if splash.done or splash.frame.isNull() or not splash.grab().save(path):
                app.exit(3)
        QTimer.singleShot(1600, capture_splash)
    return app.exec_()


if __name__ == '__main__':
    sys.exit(main())
